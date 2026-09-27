"""Render deliverables/investment_note.md to deliverables/investment_note.docx (python-docx). Minimal markdown: #, ##, tables, *italic*, **bold**."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = Path(__file__).resolve().parents[1] / "deliverables" / "investment_note.md"
DST = SRC.with_suffix(".docx")


def add_runs(par, text):
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*"):
            par.add_run(tok[1:-1]).italic = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
        else:
            par.add_run(tok)


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.6); s.left_margin = s.right_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(9.5)
    st.paragraph_format.space_after = Pt(4)
    lines = SRC.read_text().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "):
            h = doc.add_heading(ln[2:], level=1); h.runs[0].font.size = Pt(15)
        elif ln.startswith("## "):
            h = doc.add_heading(ln[3:], level=2); h.runs[0].font.size = Pt(11.5)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|\s*-", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    p = t.cell(r, c).paragraphs[0]; add_runs(p, cell)
                    for run in p.runs: run.font.size = Pt(8.5)
            doc.add_paragraph()
            continue
        elif ln.strip():
            p = doc.add_paragraph(); add_runs(p, ln); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1
    doc.save(DST); print("wrote", DST)


if __name__ == "__main__":
    main()
